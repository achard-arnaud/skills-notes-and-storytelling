#!/usr/bin/env python3
import argparse
import re
from pathlib import Path

FR_JARGON = {
    "harvest": "adaptation / capitalisation / inspiration",
    "catch-up": "rattrapage",
    "catch up": "rattrapage",
    "speed-up": "accélération",
    "speed up": "accélération",
    "hard gate": "critère bloquant",
    "buyer": "acheteur / client",
    "seller": "vendeur / fournisseur",
}
EN_JARGON = {"roadmap de": "roadmap / implementation plan"}

FR_STOP = {"le","la","les","un","une","des","de","du","dans","pour","avec","sur","est","sont","et","ou","que","qui","au","aux","ce","cette","ces","plus","comme","par","en","à","afin","sans"}
EN_STOP = {"the","a","an","of","to","in","for","with","on","is","are","and","or","that","which","this","these","from","as","by","into","without","can","will","should"}

URL_RE = re.compile(r"https?://\S+")
CODE_RE = re.compile(r"\x60[^\x60]+\x60")
WORD_RE = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ'-]+")

def clean(text):
    return CODE_RE.sub(" ", URL_RE.sub(" ", text))

def paragraphs(text):
    return [p.strip() for p in re.split(r"\n\s*\n|(?<=\.)\s+(?=[A-ZÀ-Ö])", text) if p.strip()]

def word_counts(p):
    words=[w.lower() for w in WORD_RE.findall(clean(p))]
    return words, sum(w in FR_STOP for w in words), sum(w in EN_STOP for w in words)

def lint_text(text, language, allow_terms):
    issues=[]
    low=clean(text).lower()
    allow={x.lower() for x in allow_terms}
    mapping=FR_JARGON if language.startswith("fr") else EN_JARGON if language.startswith("en") else {}
    for term,repl in mapping.items():
        if term in low and term not in allow:
            prefix="préférer" if language.startswith("fr") else "prefer"
            issues.append(f"jargon:{term} -> {prefix} {repl}")
    for i,p in enumerate(paragraphs(text),1):
        words,fr,en=word_counts(p)
        if len(words) < 12:
            continue
        if language.startswith("fr") and en >= 5 and en > max(2, fr*2):
            issues.append(f"language-drift: paragraphe {i} paraît majoritairement anglais")
        elif language.startswith("en") and fr >= 5 and fr > max(2, en*2):
            issues.append(f"language-drift: paragraph {i} appears predominantly French")
    return issues

def read_docx(path):
    try:
        from docx import Document
    except ImportError as e:
        raise SystemExit("python-docx is required for --docx") from e
    d=Document(path)
    chunks=[p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            chunks.extend(c.text for c in r.cells)
    return "\n\n".join(chunks)

def self_test():
    bad="Le catch-up produit repose sur un harvest des patterns et un hard gate commercial."
    good="Le rattrapage produit repose sur une capitalisation des modèles et un critère bloquant commercial."
    assert lint_text(bad,"fr",[])
    assert not lint_text(good,"fr",[])
    drift="The platform is designed for enterprise teams and the system will coordinate agents with the workflow and the tools."
    assert any("language-drift" in x for x in lint_text(drift,"fr",[]))
    print("language-lint: PASS")

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--language")
    group=ap.add_mutually_exclusive_group()
    group.add_argument("--file")
    group.add_argument("--docx")
    group.add_argument("--text")
    ap.add_argument("--allow-term", action="append", default=[])
    ap.add_argument("--self-test", action="store_true")
    args=ap.parse_args()
    if args.self_test:
        self_test(); return
    if not args.language:
        raise SystemExit("--language is required")
    if args.file:
        text=Path(args.file).read_text(encoding="utf-8")
    elif args.docx:
        text=read_docx(args.docx)
    elif args.text:
        text=args.text
    else:
        raise SystemExit("provide --file, --docx or --text")
    issues=lint_text(text,args.language,args.allow_term)
    if issues:
        for issue in issues:
            print(issue)
        raise SystemExit(1)
    print("language-lint: PASS")

if __name__=="__main__":
    main()
